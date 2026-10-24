# -*- coding: utf-8 -*-
"""
assets.py —— 存量用例资产读取器（把 user 手上的历史 Excel 变成可用的结构化数据）

为什么不能用 openpyxl 默认方式直接读
    user 的历史用例分三代、格式各不相同：
      ① 归档版 / 930母版        A–P 16 列
      ② 早期特性用例            12 列（无截图列）
      ③ 开发冒烟自测 / 含结果版  23–24 列，且**带 dataValidation 的新字段**，
                                openpyxl 3.1.5 直接 load 会 TypeError
                                → 必须 `read_only=True` 退回只读模式
    另外这些文件的列顺序/列名不统一（有的带"（可为空）"后缀），所以一律**按列名定位**，
    不按列号。

读出来的用途（对应 user 的诉求）
    · 覆盖地图：某个功能模块下已有哪些场景 → 生成新用例前先查重、补缺
    · 偏差清单：`测试结果` 标了失败/不通过的用例 → 就是 user 说的"与实际效果不符"
      的那批，正好当作"待验证假设"而不是事实
    · 环境事实抽取：从"前置条件/测试数据"里捞出项目名、场次、票档、价格 → 交给
      catalog/探针去实测核对

★ 权威源（2026-09-30 user 定）
    同一特性往往有 3 份文件：纯用例版 / **（含结果）版** / -开发冒烟自测版。
    **只采信「含结果」版**，同族的其余文件一律丢弃：
      · `scan(only="含结果")`（CLI 默认）→ 930 目录只有 5 份 = 267 条用例 / 578 步骤行；
      · `scan(authoritative=True)`（默认）→ 同族自动留「含结果」版，没有含结果版的文件保留。
    别把 16 份文件混在一起数 —— 那是虚高数字。
"""
from __future__ import annotations

import collections
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from . import xlcompat  # noqa: F401  —— 打兼容补丁，必须在 load_workbook 之前 import

# 表头归一化：去掉"（可为空）"这类后缀与空白，统一到标准列名
_ALIASES = {
    "用例编号": "case_id", "编号": "case_id", "用例id": "case_id",
    "优先级": "priority",
    "功能模块": "module", "模块": "module",
    "用例标题": "title", "标题": "title",
    "前置条件": "precondition",
    "操作步骤": "action", "步骤": "action",
    "预期结果": "expect",
    "测试数据": "data",
    "测试结果": "result",
    "可自动化": "auto",
    "开发自测": "dev_self", "自测结果": "dev_self",
    "备注": "note",
}
_FAIL_WORDS = ("失败", "不通过", "未通过", "异常", "阻塞", "fail", "ng", "×", "✗")

# 权威版本标记：同一特性的多份文件里，带「含结果」的那份是**唯一权威源**
# （2026-09-30 user 定：以含结果为准；纯用例版里的差异一律不采信）
AUTH_MARK = "含结果"

# 同族归一：把「多次核销测试用例（含结果）」「多次核销测试用例」「多次核销-开发冒烟自测(x)」
# 都收敛到同一个族名，这样才知道"这几份是同一件事的不同版本"，从而只留权威那份。
_FAMILY_STRIP = (
    r"[-_—\s]*开发冒烟自测.*$",
    r"[-_—\s]*执行结果.*$",
    r"[-_—\s]*测试用例$",
    r"[-_—\s]*用例$",
)


def file_family(name: str) -> str:
    """文件名 → 族名（去括号后缀、去"开发冒烟自测/执行结果/测试用例"等版本尾巴）。"""
    stem = re.sub(r"[（(][^）)]*[）)]", "", Path(name).stem)
    for pat in _FAMILY_STRIP:
        stem = re.sub(pat, "", stem)
    return stem.strip() or Path(name).stem


def is_authoritative(path: str | Path) -> bool:
    return AUTH_MARK in Path(path).name


def _norm_header(text) -> str | None:
    if text is None:
        return None
    t = re.sub(r"[（(][^）)]*[）)]", "", str(text)).strip().replace(" ", "")
    return _ALIASES.get(t)


def _norm_result(text) -> str:
    t = str(text or "").strip()
    if not t:
        return "blank"
    low = t.lower()
    if any(w in low for w in _FAIL_WORDS):
        return "fail"
    if any(w in low for w in ("通过", "pass", "ok", "√", "✓")):
        return "pass"
    return "other"


@dataclass
class RefCase:
    case_id: str
    priority: str = ""
    module: str = ""
    title: str = ""
    precondition: str = ""
    steps: list[tuple[str, str]] = field(default_factory=list)   # (操作步骤, 预期结果)
    result: str = ""
    result_kind: str = "blank"
    auto: str = ""
    dev_self: str = ""
    note: str = ""

    @property
    def n_steps(self) -> int:
        return len(self.steps)


@dataclass
class FileAssets:
    path: Path
    sheet: str = ""
    n_cols: int = 0
    headers: list[str] = field(default_factory=list)
    cases: list[RefCase] = field(default_factory=list)
    error: str = ""
    mode: str = ""            # normal / read_only

    @property
    def modules(self) -> dict[str, int]:
        return dict(collections.Counter(c.module for c in self.cases if c.module))

    @property
    def fails(self) -> list[RefCase]:
        return [c for c in self.cases if c.result_kind == "fail"]

    @property
    def n_steps(self) -> int:
        return sum(c.n_steps for c in self.cases)


def _load(path: Path):
    """返回 (ws, mode)。含结果版会退回只读模式。"""
    try:
        return openpyxl.load_workbook(path).active, "normal"
    except TypeError:
        return openpyxl.load_workbook(path, read_only=True).active, "read_only"


def read_cases(path: str | Path) -> FileAssets:
    path = Path(path)
    fa = FileAssets(path=path)
    try:
        ws, mode = _load(path)
    except Exception as e:
        fa.error = f"{type(e).__name__}: {e}"
        return fa
    fa.mode = mode
    try:
        it = ws.iter_rows(values_only=True)
        header = next(it)
    except StopIteration:
        fa.error = "空表"
        return fa

    fa.sheet = ws.title
    fa.n_cols = len(header)
    fa.headers = [str(h) if h is not None else "" for h in header]
    col = {}
    for i, h in enumerate(header):
        k = _norm_header(h)
        if k and k not in col:
            col[k] = i
    if "case_id" not in col or "action" not in col:
        fa.error = f"表头不像用例表（缺 用例编号/操作步骤）：{fa.headers[:8]}"
        return fa

    cur: RefCase | None = None
    for row in it:
        get = lambda k: (str(row[col[k]]).strip() if k in col and row[col[k]] is not None
                         else "")
        cid = get("case_id")
        if cid and cid not in ("None",):
            cur = RefCase(case_id=cid)
            cur.priority = get("priority")
            cur.module = get("module")
            cur.title = get("title")
            cur.precondition = get("precondition")
            cur.result = get("result")
            cur.result_kind = _norm_result(cur.result)
            cur.auto = get("auto")
            cur.dev_self = get("dev_self")
            cur.note = get("note")
            fa.cases.append(cur)
        if cur is None:
            continue
        action, expect = get("action"), get("expect")
        if action or expect:
            cur.steps.append((action, expect))
        # 单步用例：结果可能写在有结果的那一行（合并前）
        if not cur.result:
            r = get("result")
            if r:
                cur.result, cur.result_kind = r, _norm_result(r)
    return fa


def _candidates(root: Path, glob: str, max_mb: float, skip: tuple[str, ...],
                only: str | None) -> tuple[list[Path], list[tuple[Path, str]]]:
    """挑出可读的候选文件，并记录被筛掉的原因。"""
    files, dropped = [], []
    for p in sorted(root.rglob(glob)):
        if p.suffix.lower() not in (".xlsx", ".xlsm") or p.name.startswith("~$"):
            continue
        if max_mb and p.stat().st_size > max_mb * 1024 * 1024:
            dropped.append((p, f"体积 {p.stat().st_size / 1048576:.0f}MB > {max_mb}MB"))
            continue
        if any(s in p.name for s in skip):
            dropped.append((p, "命中 skip"))
            continue
        if only and only not in p.name:
            dropped.append((p, f"不含「{only}」"))
            continue
        files.append(p)
    return files, dropped


def pick_authoritative(files: list[Path]) -> tuple[list[Path], list[tuple[Path, Path]]]:
    """同族去重：族内有「含结果」版 → 只留它，其余记入 replaced。

    返回 (留下的文件, [(被替换的文件, 顶替它的文件)])。
    """
    by_family: dict[str, list[Path]] = collections.defaultdict(list)
    for p in files:
        by_family[file_family(p.name)].append(p)
    keep, replaced = [], []
    for _fam, group in sorted(by_family.items()):
        auth = [p for p in group if is_authoritative(p)]
        if auth:
            winner = sorted(auth)[0]
            keep.append(winner)
            replaced += [(p, winner) for p in group if p != winner]
        else:
            keep.extend(group)
    return sorted(keep), replaced


def scan(root: str | Path, glob: str = "*.xlsx", max_mb: float = 8.0,
         skip: tuple[str, ...] = (), only: str | None = None,
         authoritative: bool = True) -> list[FileAssets]:
    """读取目录下的存量用例。

    authoritative=True（默认）：同族只留「含结果」版 —— 它是唯一权威源。
    only：只读文件名含该子串的文件（例如 only="含结果"）。
    """
    root = Path(root)
    files, _ = _candidates(root, glob, max_mb, skip, only)
    if authoritative:
        files, _ = pick_authoritative(files)
    return [read_cases(p) for p in files]


def scan_report(root: str | Path, glob: str = "*.xlsx", max_mb: float = 8.0,
                skip: tuple[str, ...] = (), only: str | None = None,
                authoritative: bool = True):
    """scan 的带说明版：返回 (assets, dropped, replaced)，便于打印"为什么没读那几份"。"""
    root = Path(root)
    files, dropped = _candidates(root, glob, max_mb, skip, only)
    replaced: list[tuple[Path, Path]] = []
    if authoritative:
        files, replaced = pick_authoritative(files)
    return [read_cases(p) for p in files], dropped, replaced


# ------------------------------------------------------------------ 查已有用例


def find(refs: list[FileAssets], query: str = "", module: str = "",
         limit: int = 60) -> list[tuple[FileAssets, RefCase]]:
    """在存量资产里按关键词查用例（生成新用例前查重用）。"""
    q = (query or "").lower()
    m = (module or "").lower()
    hit = []
    for fa in refs:
        if fa.error:
            continue
        for c in fa.cases:
            if m and m not in (c.module or "").lower():
                continue
            if q:
                blob = " ".join([c.case_id, c.title, c.module, c.precondition,
                                 c.note] + [a + " " + e for a, e in c.steps]).lower()
                if q not in blob:
                    continue
            hit.append((fa, c))
            if len(hit) >= limit:
                return hit
    return hit


def corpus_stats(refs: list[FileAssets]) -> dict:
    """权威语料的规模与风格画像（写进覆盖地图，供生成用例时对齐）。"""
    ok = [a for a in refs if not a.error]
    per_file = []
    for fa in ok:
        multi = [c for c in fa.cases if c.n_steps > 1]
        all_step = sum(1 for c in multi if all(e.strip() for _, e in c.steps))
        per_file.append({
            "file": fa.path.name,
            "sheet": fa.sheet,
            "cols": fa.n_cols,
            "cases": len(fa.cases),
            "steps": fa.n_steps,
            "modules": len(fa.modules),
            "fails": [c.case_id for c in fa.fails],
            "style": {
                "multi_step_cases": len(multi),
                "every_step_has_expect": all_step,
                "expect_last_step_only": len(multi) - all_step,
            },
        })
    multi_all = sum(p["style"]["multi_step_cases"] for p in per_file)
    every_all = sum(p["style"]["every_step_has_expect"] for p in per_file)
    return {
        "authoritative": True,
        "mark": AUTH_MARK,
        "files": len(ok),
        "cases": sum(len(a.cases) for a in ok),
        "steps": sum(a.n_steps for a in ok),
        "modules": len({c.module for a in ok for c in a.cases if c.module}),
        "fails": sum(len(a.fails) for a in ok),
        "multi_step_cases": multi_all,
        "every_step_has_expect": every_all,
        "per_file": per_file,
    }



def summarize(assets: list[FileAssets]) -> str:
    lines = [f"{'文件':<46}{'列':>4}{'用例':>5}{'步骤':>5}  {'模块数':>5} {'失败':>4} 模式"]
    lines.append("-" * 92)
    for fa in assets:
        name = fa.path.name if len(fa.path.name) <= 44 else fa.path.name[:41] + "…"
        if fa.error:
            lines.append(f"{name:<46}{'—':>4}{'—':>5}{'—':>5}  {'—':>5} {'—':>4} ✗ {fa.error[:40]}")
            continue
        lines.append(f"{name:<46}{fa.n_cols:>4}{len(fa.cases):>5}{fa.n_steps:>5}  "
                     f"{len(fa.modules):>5} {len(fa.fails):>4} {fa.mode}")
    ok = [a for a in assets if not a.error]
    lines.append("-" * 92)
    lines.append(f"共 {len(assets)} 份（可读 {len(ok)}）：用例 "
                 f"{sum(len(a.cases) for a in ok)} 条 / 步骤行 {sum(a.n_steps for a in ok)} 行 / "
                 f"标为失败 {sum(len(a.fails) for a in ok)} 条")
    return "\n".join(lines)


def build_index(assets: list[FileAssets], out: str | Path, root: str | Path | None = None,
                dropped: list | None = None, replaced: list | None = None) -> Path:
    """写覆盖地图：模块 → 用例（含来源文件、结果状态）+ 语料统计。"""
    modules: dict[str, list[dict]] = collections.defaultdict(list)
    files = []
    for fa in assets:
        if fa.error:
            files.append({"file": str(fa.path), "error": fa.error})
            continue
        files.append({
            "file": str(fa.path), "rel": fa.path.name, "sheet": fa.sheet, "cols": fa.n_cols,
            "cases": len(fa.cases), "steps": fa.n_steps,
            "modules": fa.modules, "mode": fa.mode,
            "authoritative": is_authoritative(fa.path),
            "fails": [c.case_id for c in fa.fails],
        })
        for c in fa.cases:
            modules[c.module or "(未填模块)"].append({
                "case_id": c.case_id, "title": c.title, "priority": c.priority,
                "steps": c.n_steps, "result": c.result_kind,
                "source": fa.path.name,
            })
    doc = {
        "root": str(root) if root else None,
        "policy": {
            "authoritative_mark": AUTH_MARK,
            "note": "同族文件只采信带「含结果」的那份；纯用例版的差异不采信。",
            "dropped": [{"file": str(p), "why": w} for p, w in (dropped or [])],
            "replaced": [{"dropped": str(a), "kept": str(b)} for a, b in (replaced or [])],
        },
        "stats": corpus_stats(assets),
        "files": files,
        "modules": {k: {"n": len(v), "cases": v} for k, v in sorted(modules.items())},
    }
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")
    return out
